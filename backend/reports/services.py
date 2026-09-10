import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from customers.models import Customer
from orders.models import Order
from groceries.models import GroceryItem
from feedback.models import Feedback

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_FILE_PATH = BASE_DIR / 'kanakis_store_orders.xlsx'

def generate_master_excel(output_path=None, filter_type=None, customer_id=None, from_date=None, to_date=None):
    if output_path is None:
        output_path = EXCEL_FILE_PATH

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    forest_green_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E0DCD3'),
        right=Side(style='thin', color='E0DCD3'),
        top=Side(style='thin', color='E0DCD3'),
        bottom=Side(style='thin', color='E0DCD3')
    )
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. Customers Sheet
    ws_customers = wb.create_sheet(title="Customers")
    headers_customers = ["Customer ID", "Customer Name", "Mobile", "Email", "Address", "City", "Pincode", "Registration Date"]
    ws_customers.append(headers_customers)

    customers_qs = Customer.objects.all()
    for c in customers_qs:
        ws_customers.append([
            f"CUST-{c.id:03d}",
            c.name,
            c.mobile,
            c.email or "—",
            c.address,
            c.city,
            c.pincode or "—",
            c.created_at.strftime("%Y-%m-%d")
        ])

    # 2. Orders Sheet
    ws_orders = wb.create_sheet(title="Orders")
    headers_orders = ["Order ID", "Customer ID", "Customer Name", "Mobile", "Order Date", "Total Items", "Source", "Status"]
    ws_orders.append(headers_orders)

    orders_qs = Order.objects.select_related('customer').prefetch_related('items').all()
    if customer_id:
        orders_qs = orders_qs.filter(customer__id=customer_id)
    if from_date and to_date:
        orders_qs = orders_qs.filter(order_date__range=[from_date, to_date])

    for o in orders_qs:
        ws_orders.append([
            o.order_number,
            f"CUST-{o.customer.id:03d}",
            o.customer.name,
            o.customer.mobile,
            o.order_date.strftime("%Y-%m-%d"),
            o.items.count(),
            o.source,
            o.status
        ])

    # 3. Grocery Items Sheet
    ws_items = wb.create_sheet(title="Grocery Items")
    headers_items = ["Order ID", "Customer Name", "Grocery Item", "Quantity", "Unit", "Notes"]
    ws_items.append(headers_items)

    for o in orders_qs:
        for item in o.items.all():
            ws_items.append([
                o.order_number,
                o.customer.name,
                item.name,
                float(item.quantity),
                item.unit,
                item.notes or "—"
            ])

    # 4. Feedback Sheet
    ws_feedback = wb.create_sheet(title="Feedback")
    headers_feedback = ["Feedback ID", "Customer Name", "Mobile", "Rating", "Feedback", "Date"]
    ws_feedback.append(headers_feedback)

    feedback_qs = Feedback.objects.all()
    for fb in feedback_qs:
        ws_feedback.append([
            f"FB-{fb.id:03d}",
            fb.customer_name,
            fb.mobile,
            f"{fb.rating} ★",
            fb.message,
            fb.created_at.strftime("%Y-%m-%d")
        ])

    # Format all sheets
    for ws in [ws_customers, ws_orders, ws_items, ws_feedback]:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 26

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = forest_green_fill
            cell.font = header_font
            cell.alignment = header_align

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return output_path
